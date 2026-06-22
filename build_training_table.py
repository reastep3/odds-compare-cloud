"""
Stage 0 of the auto-calibration workstream: build ONE pick-level training
table from the graded backtest reports, then print a baseline that tells us
whether the dataset is rich enough to fit a meta-model on (the model that
would eventually replace the hand-tuned gates in odds_compare._qualifies /
_qualifies_prop).

This is ETL, not ML. It does three things:

  1. Reads every graded report (backtest_*.xlsx, _bt_recent/, Backtest
     Archives/) — the same file set analyze_ev_tiers.py uses — pulling the
     BROAD +EV pool from "All Picks" (game lines) and "Player Props". The
     "Promo P&L" sheet is intentionally excluded: promo money is play money
     and never rolls into the main P&L (see project rules).

  2. Dedups the heavy cross-file overlap into one canonical row per pick,
     merging toward completeness (settled result + captured CLV win).

  3. Writes training_table.csv (one row per pick, features + labels) and
     prints a baseline report: label coverage, per-(sport x market) sample
     richness, and a CLV-validity check — does a positive closing-line value
     actually predict positive realized results in *our* data? That last
     check is the whole premise of using CLV as the fast label, so we test
     it before investing in a model.

Two labels per pick:
    realized P&L / result  — high variance, needs hundreds of bets
    clv_pct                — low variance, needs dozens (game lines only;
                             props carry no CLV in the reports yet)

Run:  python build_training_table.py
"""
import csv
import glob
import os
import re
from collections import defaultdict
from datetime import date

import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT_CSV = os.path.join(ROOT, "training_table.csv")

SETTLED = ("WIN", "LOSS", "PUSH")


# ───────────────────────── parsing helpers ──────────────────────────

def parse_american(v):
    """'+128' / '-110' -> int, or None if unparseable."""
    if v is None:
        return None
    m = re.match(r"^([+-]?)(\d+)$", str(v).strip().replace(" ", ""))
    if not m:
        return None
    sign, num = m.groups()
    return -int(num) if sign == "-" else int(num)


def american_to_decimal(a):
    return 1 + (a / 100 if a > 0 else 100 / -a)


def implied_prob(a):
    """American odds -> raw implied probability (vig included)."""
    if a is None:
        return None
    return (-a) / ((-a) + 100) if a < 0 else 100 / (a + 100)


def parse_date_cell(v):
    if isinstance(v, date):
        return v
    if v is None:
        return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", str(v).strip())
    if not m:
        return None
    try:
        return date(int(m[1]), int(m[2]), int(m[3]))
    except ValueError:
        return None


def to_float(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ───────────────────────── file discovery ───────────────────────────

def all_files():
    files = sorted(glob.glob(os.path.join(ROOT, "backtest_*.xlsx")))
    files += sorted(glob.glob(os.path.join(ROOT, "_bt_recent", "*.xlsx")))
    files += sorted(glob.glob(os.path.join(ROOT, "Backtest Archives", "*.xlsx")))
    return files


# ───────────────────────── row extraction ───────────────────────────

def _row_record(r, ix, is_prop):
    """Turn one sheet row into a normalized record dict, or None to skip."""
    game_date = parse_date_cell(r[ix["Game Date"]]) if "Game Date" in ix else None
    if game_date is None:
        return None
    odds = parse_american(r[ix["Odds"]]) if "Odds" in ix else None
    if odds is None:
        return None

    sport = str(r[ix["Sport"]]).upper() if "Sport" in ix else ""
    game = r[ix["Game"]] if "Game" in ix else ""
    book = r[ix["Book"]] if "Book" in ix else ""
    market = str(r[ix["Market"]]).lower() if "Market" in ix else ""
    line = r[ix["Line"]] if "Line" in ix else None
    result = r[ix["Result"]] if "Result" in ix else None
    if result not in SETTLED:
        result = None  # NO_RESULT / blank -> unsettled

    if is_prop:
        player = r[ix["Player"]] if "Player" in ix else ""
        side = r[ix["Side"]] if "Side" in ix else ""
        pick = f"{player}/{side}"
    else:
        player = side = ""
        pick = r[ix["Pick"]] if "Pick" in ix else ""

    bet = to_float(r[ix["Bet Amt"]]) if "Bet Amt" in ix else None
    pnl = to_float(r[ix["P&L"]]) if "P&L" in ix else None
    if pnl is None and result is not None:
        # Reconstruct realized P&L from the result when the cell is blank
        # (mirrors analyze_ev_tiers.load).
        b = bet or 0.0
        dec = american_to_decimal(odds)
        pnl = (dec - 1) * b if result == "WIN" else (-b if result == "LOSS" else 0.0)

    # Flat unit P&L: each pick = 1u staked (robust to Kelly outliers).
    if result == "WIN":
        unit_pnl = american_to_decimal(odds) - 1
    elif result == "LOSS":
        unit_pnl = -1.0
    else:
        unit_pnl = 0.0  # PUSH or unsettled

    file_date = parse_date_cell(r[ix["File Date"]]) if "File Date" in ix else None
    clv = to_float(r[ix["CLV %"]]) if "CLV %" in ix else None
    close = r[ix["Close"]] if "Close" in ix else None
    ev = to_float(r[ix["EV %"]]) if "EV %" in ix else None
    confidence = r[ix["Confidence"]] if "Confidence" in ix else None
    seen = to_float(r[ix["Seen"]]) if "Seen" in ix else None
    persistent = r[ix["Persistent"]] if "Persistent" in ix else None

    return {
        "is_prop": is_prop,
        "game_date": game_date,
        "file_date": file_date,
        "sport": sport,
        "game": str(game),
        "book": str(book),
        "market": market,
        "player": str(player),
        "side": str(side),
        "pick": str(pick),
        "line": "" if line is None else str(line),
        "odds": odds,
        "ev": ev,
        "confidence": "" if confidence is None else str(confidence),
        "seen": seen,
        "persistent": "" if persistent is None else str(persistent),
        "result": result,
        "bet": bet,
        "pnl": pnl,
        "unit_pnl": unit_pnl,
        "clv_pct": clv,
        "close": "" if close is None else str(close),
    }


def dedup_key(rec):
    return (rec["game_date"].isoformat(), rec["sport"], rec["game"],
            rec["book"], rec["market"], rec["pick"], rec["line"],
            rec["odds"], rec["is_prop"])


def merge(a, b):
    """Merge two records for the same pick toward completeness."""
    out = dict(a)
    a_settled = a["result"] in SETTLED
    b_settled = b["result"] in SETTLED
    # Prefer a settled result over an unsettled one.
    if not a_settled and b_settled:
        out.update(result=b["result"], pnl=b["pnl"], unit_pnl=b["unit_pnl"],
                   bet=b["bet"] if b["bet"] is not None else a["bet"])
    # Prefer a captured CLV over a missing one.
    if a.get("clv_pct") is None and b.get("clv_pct") is not None:
        out["clv_pct"] = b["clv_pct"]
        out["close"] = b["close"]
    # Earliest file_date = first time the pick was seen (for lead time).
    if b["file_date"] and (not a["file_date"] or b["file_date"] < a["file_date"]):
        out["file_date"] = b["file_date"]
    # Keep the larger 'seen' count if both present.
    if b["seen"] is not None and (a["seen"] is None or b["seen"] > a["seen"]):
        out["seen"] = b["seen"]
    # Backfill EV / confidence if a lacked them.
    for f in ("ev", "confidence", "persistent"):
        if not out.get(f) and b.get(f):
            out[f] = b[f]
    return out


def load():
    """Dedup'd broad-pool picks across every graded report."""
    bets = {}
    files = all_files()
    skipped = 0
    for fp in files:
        try:
            wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
        except Exception as e:                       # noqa: BLE001
            print(f"  skip {os.path.basename(fp)}: {e}")
            skipped += 1
            continue
        for sheet, is_prop in (("All Picks", False), ("Player Props", True)):
            if sheet not in wb.sheetnames:
                continue
            it = wb[sheet].iter_rows(values_only=True)
            hdr = next(it, None)
            if not hdr:
                continue
            ix = {n: i for i, n in enumerate(hdr) if n is not None}
            for r in it:
                try:
                    rec = _row_record(r, ix, is_prop)
                except (KeyError, IndexError, TypeError, ValueError):
                    continue
                if rec is None:
                    continue
                k = dedup_key(rec)
                bets[k] = merge(bets[k], rec) if k in bets else rec
        wb.close()
    return bets, len(files) - skipped


# ───────────────────────── CSV output ────────────────────────────────

CSV_COLS = [
    "is_prop", "game_date", "file_date", "lead_time_days", "sport", "market",
    "book", "pick", "player", "side", "line", "odds", "odds_decimal",
    "implied_prob", "ev", "confidence", "seen", "persistent",
    "result", "settled", "won", "bet", "pnl", "unit_pnl",
    "clv_pct", "clv_captured", "close",
]


def write_csv(bets):
    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(CSV_COLS)
        for rec in bets.values():
            settled = rec["result"] in SETTLED
            won = "" if rec["result"] not in ("WIN", "LOSS") else int(rec["result"] == "WIN")
            lead = ""
            if rec["file_date"] and rec["game_date"]:
                lead = (rec["game_date"] - rec["file_date"]).days
            w.writerow([
                int(rec["is_prop"]),
                rec["game_date"].isoformat(),
                rec["file_date"].isoformat() if rec["file_date"] else "",
                lead, rec["sport"], rec["market"], rec["book"], rec["pick"],
                rec["player"], rec["side"], rec["line"], rec["odds"],
                round(american_to_decimal(rec["odds"]), 4),
                round(implied_prob(rec["odds"]), 5),
                rec["ev"] if rec["ev"] is not None else "",
                rec["confidence"],
                rec["seen"] if rec["seen"] is not None else "",
                rec["persistent"],
                rec["result"] or "", int(settled), won,
                rec["bet"] if rec["bet"] is not None else "",
                round(rec["pnl"], 4) if rec["pnl"] is not None else "",
                round(rec["unit_pnl"], 4),
                rec["clv_pct"] if rec["clv_pct"] is not None else "",
                int(rec["clv_pct"] is not None),
                rec["close"],
            ])


# ───────────────────────── baseline report ──────────────────────────

def _perf(picks):
    """(n_settled, w, l, roi%, flat%, avg_clv or None)."""
    s = [p for p in picks if p["result"] in SETTLED]
    w = sum(1 for p in s if p["result"] == "WIN")
    l = sum(1 for p in s if p["result"] == "LOSS")
    wager = sum(p["bet"] or 0 for p in s)
    pnl = sum(p["pnl"] or 0 for p in s)
    roi = pnl / wager * 100 if wager else 0.0
    flat = sum(p["unit_pnl"] for p in s) / len(s) * 100 if s else 0.0
    clvs = [p["clv_pct"] for p in s if p["clv_pct"] is not None]
    avg_clv = sum(clvs) / len(clvs) if clvs else None
    return len(s), w, l, roi, flat, avg_clv


def _line(label, picks, indent=0):
    n, w, l, roi, flat, avg_clv = _perf(picks)
    if n == 0:
        return
    clv = f"  CLV {avg_clv:+5.2f}%" if avg_clv is not None else "  CLV   n/a"
    print(f"{' ' * indent}{label:<26} n={n:>4}  W{w:>4}/L{l:<4}  "
          f"ROI {roi:>+6.1f}%  flat {flat:>+6.1f}%{clv}")


def section(t):
    print("\n" + "=" * 84)
    print(t)
    print("=" * 84)


def report(bets, files_read):
    vals = list(bets.values())
    games = [v for v in vals if not v["is_prop"]]
    props = [v for v in vals if v["is_prop"]]

    dates = [v["game_date"] for v in vals]
    section("DATASET")
    print(f"Files read: {files_read}")
    print(f"Unique picks: {len(vals)}   ({len(games)} game-line, {len(props)} prop)")
    print(f"Game-date span: {min(dates)} .. {max(dates)}")
    print(f"Written to: {os.path.relpath(OUT_CSV, ROOT)}")

    # ---- label coverage ----
    section("LABEL COVERAGE")
    for label, grp in (("Game lines", games), ("Props", props)):
        n = len(grp)
        settled = sum(1 for p in grp if p["result"] in SETTLED)
        clv = sum(1 for p in grp if p["clv_pct"] is not None)
        print(f"{label:<12} total={n:>5}  settled={settled:>5} "
              f"({settled / n * 100:>4.0f}%)  CLV-captured={clv:>5} "
              f"({clv / n * 100:>4.0f}%)" if n else f"{label}: none")

    # ---- overall performance ----
    section("BASELINE PERFORMANCE (broad +EV pool, pre-filter)")
    _line("ALL", vals)
    _line("Game lines", games)
    _line("Props", props)

    # ---- per-sport ----
    section("BY SPORT (settled)")
    by_sport = defaultdict(list)
    for v in vals:
        by_sport[v["sport"]].append(v)
    for sp in sorted(by_sport, key=lambda s: -_perf(by_sport[s])[0]):
        _line(sp or "(blank)", by_sport[sp])

    # ---- richness: how many (sport x market) cells clear sample bars ----
    section("SAMPLE RICHNESS — settled picks per (sport x market) cell")
    for label, grp in (("GAME LINES", games), ("PROPS", props)):
        by_cell = defaultdict(int)
        for p in grp:
            if p["result"] in SETTLED:
                by_cell[(p["sport"], p["market"])] += 1
        cells = sorted(by_cell.values(), reverse=True)
        ge = lambda t: sum(1 for c in cells if c >= t)  # noqa: E731
        print(f"{label}: {len(cells)} cells | "
              f">=30: {ge(30)}  >=100: {ge(100)}  >=200: {ge(200)}  "
              f">=500: {ge(500)}")
        for (sp, mk), c in sorted(by_cell.items(), key=lambda x: -x[1])[:8]:
            print(f"    {sp:<10} {mk:<16} n={c}")

    # ---- CLV-validity check: the core Stage-0 question ----
    section("CLV-VALIDITY CHECK (game lines, settled, CLV captured)")
    print("Does a positive closing-line value predict better realized results?")
    print("If CLV>0 outperforms CLV<=0 here, CLV is a usable fast label.\n")
    cl = [p for p in games if p["result"] in SETTLED and p["clv_pct"] is not None]
    pos = [p for p in cl if p["clv_pct"] > 0]
    nonpos = [p for p in cl if p["clv_pct"] <= 0]
    _line("CLV > 0", pos)
    _line("CLV <= 0", nonpos)
    # Finer CLV bands.
    print()
    bands = [("CLV <= -5%", lambda c: c <= -5),
             ("-5% .. 0%", lambda c: -5 < c <= 0),
             ("0% .. +5%", lambda c: 0 < c <= 5),
             ("CLV > +5%", lambda c: c > 5)]
    for lbl, fn in bands:
        _line(lbl, [p for p in cl if fn(p["clv_pct"])], indent=2)

    # ---- EV-tier signal: does our own EV predict CLV/ROI? ----
    section("EV-TIER vs CLV/ROI (game lines, settled) — pricer sanity")
    tiers = [("0-1%", lambda v: v < 1), ("1-2%", lambda v: 1 <= v < 2),
             ("2-3%", lambda v: 2 <= v < 3), ("3-5%", lambda v: 3 <= v < 5),
             ("5%+", lambda v: v >= 5)]
    gl_ev = [p for p in games if p["ev"] is not None]
    for lbl, fn in tiers:
        _line(lbl, [p for p in gl_ev if fn(p["ev"])], indent=2)


def main():
    bets, files_read = load()
    write_csv(bets)
    report(bets, files_read)


if __name__ == "__main__":
    main()
